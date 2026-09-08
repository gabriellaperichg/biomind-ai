from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
import re

import fitz


@dataclass
class ParsedPage:
    page: int
    text: str
    image_count: int = 0
    title: str = ""
    blocks: list[str] = field(default_factory=list)


@dataclass
class ParsedDocument:
    document_id: str
    source: str
    path: str
    document_type: str
    title: str
    pages: list[ParsedPage]
    metadata: dict[str, str] = field(default_factory=dict)


def _document_id(path: Path) -> str:
    return re.sub(r"[^a-z0-9]+", "_", path.stem.lower()).strip("_")


def _first_title(text: str) -> str:
    for raw in text.splitlines():
        line = " ".join(raw.split()).strip()
        if not line:
            continue
        if len(line) <= 180:
            return line
    return ""


def parse_pdf(path: str | Path) -> ParsedDocument:
    path = Path(path)
    pages: list[ParsedPage] = []

    with fitz.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            text = page.get_text("text") or ""
            blocks = []
            try:
                raw_blocks = page.get_text("blocks")
                for block in raw_blocks:
                    if len(block) >= 5 and str(block[4]).strip():
                        blocks.append(" ".join(str(block[4]).split()))
            except Exception:
                blocks = []

            image_count = len(page.get_images(full=True))
            pages.append(
                ParsedPage(
                    page=page_number,
                    text=text,
                    image_count=image_count,
                    title=_first_title(text),
                    blocks=blocks,
                )
            )

    sample = "\n".join(page.text[:1500] for page in pages[:5])
    from ingestion.router import detect_document_type
    doc_type = detect_document_type(
        path,
        sample,
        len(pages),
        sum(p.image_count > 0 for p in pages),
    )

    title = _first_title(sample) or path.stem
    return ParsedDocument(
        document_id=_document_id(path),
        source=path.name,
        path=str(path),
        document_type=doc_type,
        title=title,
        pages=pages,
    )


def parse_docx(path: str | Path) -> ParsedDocument:
    from docx import Document

    path = Path(path)
    doc = Document(path)
    paragraphs: list[str] = []
    sections: list[tuple[str, str]] = []

    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split()).strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style else ""
        if "heading" in style:
            sections.append((text, style))
        paragraphs.append(text)

    # Word não possui "página" confiável sem renderização; mantemos uma unidade
    # documental única e preservamos os títulos detectados.
    text = "\n".join(paragraphs)
    doc_type = "pop" if "procedimento operacional padrão" in text.lower() else "docx"
    page = ParsedPage(page=1, text=text, title=path.stem, blocks=paragraphs)

    return ParsedDocument(
        document_id=_document_id(path),
        source=path.name,
        path=str(path),
        document_type=doc_type,
        title=path.stem,
        pages=[page],
        metadata={"heading_count": str(len(sections))},
    )


def parse_path(path: str | Path) -> ParsedDocument:
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(path)
    if suffix == ".docx":
        return parse_docx(path)
    raise ValueError(f"Formato não suportado: {path}")
